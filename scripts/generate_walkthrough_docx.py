# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

doc.add_heading("AI-Powered Stock & Signal Platform", 0)
p = doc.add_paragraph()
run = p.add_run("Complete Codebase Architecture Walkthrough")
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph(
    "This document explains how the entire trading intelligence system works: "
    "data flow, the four running services, and the exact purpose of every core file. "
    "It is based on the live codebase in D:\\stock."
)

doc.add_heading("1. Global Architecture and Data Flow", level=1)
doc.add_paragraph(
    "The platform is four subsystems running at the same time:"
)
for item in [
    "Port 8501: Streamlit frontend (0_Overview.py)",
    "Port 8000: ML Signal REST API (signals/api.py)",
    "Port 8005: Alerts and email engine (alerts/main.py)",
    "Port 8002: Backtest metrics API (backtesting/main.py)",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph(
    "Data flow: User requests a ticker in the browser. The frontend loads historical "
    "OHLCV from Yahoo Finance (yfinance), using .NS for NSE and .BO for BSE. That data "
    "goes to the ML API, which returns BUY / HOLD / SELL. The backtesting API then "
    "simulates past profitability. The dashboard charts the result. If the user created "
    "an alert, the alerts API emails a formatted HTML report via Brevo."
)

doc.add_heading("2. User Interface (Streamlit)", level=1)

doc.add_heading("0_Overview.py", level=2)
doc.add_paragraph(
    "Main website entry. Sets UTF-8 on Windows so emoji print statements do not crash "
    "the console. Calls st.set_page_config for title and wide layout. Adds the project "
    "root to sys.path so imports like from ml... work. Calls utils.api_starter."
    "ensure_api_running() so FastAPI starts with the dashboard. Defines FEATURES cards "
    "(Dashboard, Dataset, Strategy, Alerts) and renders them as HTML."
)

doc.add_heading("pages/ directory", level=2)
doc.add_paragraph(
    "1_AI_Signals.py: Live BUY/HOLD/SELL page. Sidebar ticker, period, interval. "
    "Loads prices via data/fetcher.py DataEngine.fetch_data(). Calls ML (ml/predictor.py "
    "or API on port 8000). Draws candlesticks, RSI, and the prediction card."
)
doc.add_paragraph(
    "2_Strategy_Analysis.py: Runs BacktestEngine.run_backtest(stock_data). Charts equity "
    "curve, trades, metrics. Export JSON via ui/components/export.py."
)
doc.add_paragraph(
    "3_Alerts_Preferences.py: Talks to http://localhost:8005. Create alert, instant "
    "report, stop alert, list jobs."
)
doc.add_paragraph(
    "4_Dataset_Analysis.py: Upload CSV/Excel. auto_detect_columns() maps Date, Open, "
    "High, Low, Close, Volume even when names vary (Close Price, LTP, Adj Close, etc.). "
    "Builds StockData, runs ML, shows BUY/HOLD/SELL."
)

doc.add_heading("ui/components/", level=2)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "File"
hdr[1].text = "Job"
for a, b in [
    ("charts.py", "Candlestick, RSI, MACD plots"),
    ("metrics.py", "Win rate, drawdown, Sharpe cards"),
    ("prediction_card.py", "BUY/SELL/HOLD badge"),
    ("controls.py", "Ticker / period widgets"),
    ("export.py", "Download JSON; handles missing ML signal and Metrics object"),
    ("header.py, indicators.py", "Layout and indicator tables"),
]:
    row = table.add_row().cells
    row[0].text = a
    row[1].text = b

doc.add_paragraph("ui/utils/design.py: CSS / glassmorphism theme.")
doc.add_paragraph("ui/utils/constants.py: ticker lists, session keys.")

doc.add_heading("3. Shared Contracts (contracts/schema.py)", level=1)
doc.add_paragraph(
    "Pydantic models so UI, API, and ML speak the same language."
)
doc.add_paragraph("StockData", style="Heading 3")
doc.add_paragraph(
    "Identity: symbol, current_price, price_change, market_status. "
    "OHLCV lists: dates, opens, highs, lows, closes, volumes. "
    "Indicators: rsi, sma_20, sma_50, ema_12, ema_26, macd, macd_signal, macd_hist."
)
doc.add_paragraph("MLSignal", style="Heading 3")
doc.add_paragraph(
    "action: BUY | HOLD | SELL. signal_value: 1 / 0 / -1. "
    "confidence, confidence_level, reasoning, key_factors."
)
doc.add_paragraph("BacktestMetrics / StrategyConfig", style="Heading 3")
doc.add_paragraph(
    "Capital, fees, trades, win rate, drawdown, Sharpe, comparison vs market."
)

doc.add_heading("4. Data and Market Connectors", level=1)

doc.add_heading("data/fetcher.py", level=2)
doc.add_paragraph(
    "Main market fetch. Priority: (1) local API data/api_client.py to port 8000 "
    "/supabase/... (2) pipeline if configured (3) Yahoo Finance yfinance. "
    "yfinance is what actually works for NSE/BSE (RELIANCE.NS, TCS.BO). "
    "Then it computes RSI, SMA, EMA, MACD and returns StockData."
)

doc.add_heading("data/api_client.py", level=2)
doc.add_paragraph(
    "HTTP wrapper: get_recent_data, get_ticker_data, get_latest_market, RSI search."
)

doc.add_heading("app/data_loader.py and app/data_loader_vbt.py", level=2)
doc.add_paragraph(
    "Load CSV cache or yfinance into a DataFrame with Open, High, Low, Close, Volume, "
    "Signal for VectorBT."
)

doc.add_heading("market_data/ package", level=2)
md = doc.add_table(rows=1, cols=2)
md.style = "Table Grid"
md.rows[0].cells[0].text = "File"
md.rows[0].cells[1].text = "Role"
for a, b in [
    ("nse_unofficial.py / nse_jugaad.py", "NSE website APIs. Blocked with HTTP 403 from this network."),
    ("bse_unofficial.py", "BSE endpoints. Returns HTML, not JSON."),
    ("yfinance_client.py", "Working live quotes via Yahoo Finance."),
    ("normalizer.py", "Convert any source to {exchange, symbol, last_price, ...}."),
    ("market_service.py", "One get_quote(NSE, RELIANCE) API."),
    ("signal_enhancer.py", "Extra BUY/HOLD/SELL score from RSI, MACD, 52-week, analysts."),
]:
    row = md.add_row().cells
    row[0].text = a
    row[1].text = b

doc.add_heading("5. Machine Learning Module", level=1)

doc.add_heading("signals/api.py", level=2)
doc.add_paragraph(
    "FastAPI on port 8000. Loads ml/models/rf_model.pkl and xgb_model.pkl. "
    "Features: Daily_Return, Volatility, SMA_ratio, EMA_ratio, MACD. "
    "POST /api/v1/ml/signal/live is used by alerts. Fetches history with yfinance, "
    "builds features, ensemble vote to BUY/SELL/HOLD."
)

doc.add_heading("ml/predictor.py", level=2)
doc.add_paragraph(
    "MLEngine used by Streamlit. Loads RF + XGBoost. Optional LSTM (pradict_lstm.py). "
    "Builds MLSignal with reasoning and confidence."
)

doc.add_heading("ml/genai_ensem.py and ml/input_api.py", level=2)
doc.add_paragraph("GenAI explanation layer and extra API helpers.")

doc.add_heading("signals/train_and_save.py and signals/train_lstm.py", level=2)
doc.add_paragraph("Train RF/XGB/LSTM and write files into ml/models/.")

doc.add_heading("signals/data_pipeline.py", level=2)
doc.add_paragraph("OHLCV to feature table for training.")

doc.add_heading("6. VectorBT Simulator", level=1)

doc.add_heading("backtesting/engine_vectorbt.py", level=2)
doc.add_paragraph(
    "Class BacktestEngineVBT. INITIAL_CAPITAL = 1,000,000. FEES = 0.002 (0.2% per trade). "
    "TRADING_DAYS = 252 for CAGR and Sharpe."
)
doc.add_paragraph(
    "to_py(val): numpy scalar to Python float so JSON serialization is safe."
)
doc.add_paragraph(
    "run_market(): buy-and-hold. Daily returns become an equity curve. Computes CAGR, "
    "volatility, Sharpe, max drawdown."
)
doc.add_paragraph(
    "run_ml(): Signal == 1 enter, Signal == -1 exit. vbt.Portfolio.from_signals(...). "
    "Splits winning/losing trades, profit factor, average win/loss. Returns ml_metrics, "
    "trading_metrics, equity, trade list."
)

doc.add_heading("backtesting/engine.py", level=2)
doc.add_paragraph(
    "Same idea; wraps results in a Metrics object (ml_metrics, market_metrics, "
    "equity_curve, trades) for the Strategy page."
)

doc.add_heading("backtesting/main.py", level=2)
doc.add_paragraph(
    "FastAPI on port 8002. POST /api/v1/backtest/run {ticker} returns a confidence score "
    "for emails. Alerts call this; if it fails they fall back to live ML confidence."
)

doc.add_heading("backtesting/schemas.py", level=2)
doc.add_paragraph("Pydantic request/response for the backtest API.")

doc.add_heading("api/backtesting_api.py", level=2)
doc.add_paragraph("Older/alternate FastAPI wrapper around the same engine.")

doc.add_heading("7. Alert Services", level=1)

doc.add_heading("alerts/main.py", level=2)
doc.add_paragraph(
    "FastAPI on port 8005. Configuration: sender email plus Brevo key from "
    "scripts/email_config.py. Scheduler: APScheduler, timezone Asia/Kolkata."
)
doc.add_paragraph("fetch_ml_signal(ticker): POST port 8000 /api/v1/ml/signal/live.")
doc.add_paragraph("fetch_backtest_result(ticker): POST port 8002 /api/v1/backtest/run.")
doc.add_paragraph(
    "send_email_alert(...): HTML email via Brevo REST (xkeysib- keys) or SMTP (xsmtpsib- keys)."
)
doc.add_paragraph(
    "check_and_alert_job: live plus backtest then email (instant or cron)."
)
doc.add_paragraph("Endpoints:")
for ep in [
    "POST /create-alert: daily cron HH:MM",
    "POST /instant-report",
    "GET /active-alerts",
    "DELETE /stop-alert/...",
    "GET /health",
]:
    doc.add_paragraph(ep, style="List Bullet")

doc.add_heading("scripts/email_config.py", level=2)
doc.add_paragraph(
    "Single place for EMAIL_PROVIDER, BREVO_API_KEY, EMAIL_SENDER."
)

doc.add_heading("scripts/send_test_alert.py and send_instant_report.py", level=2)
doc.add_paragraph("CLI tests of the same email path.")

doc.add_heading("8. Supporting Files", level=1)
sup = doc.add_table(rows=1, cols=2)
sup.style = "Table Grid"
sup.rows[0].cells[0].text = "File"
sup.rows[0].cells[1].text = "Purpose"
for a, b in [
    ("requirements.txt", "Python packages"),
    ("start_services.ps1 / run_project.ps1", "Start APIs plus Streamlit"),
    ("utils/api_starter.py", "Auto-start API from Streamlit"),
    ("scripts/supabase_client.py", "Supabase URL and publishable key"),
    ("docs/SUPABASE_SETUP.md", "How to wire the database"),
    ("test_api.py, tests/system_check.py", "Health checks"),
    ("venv/", "Installed libraries. Not application code."),
]:
    row = sup.add_row().cells
    row[0].text = a
    row[1].text = b

doc.add_heading("9. End-to-End Example: Instant Report for AAPL", level=1)
doc.add_paragraph(
    "1. UI posts {user_email, ticker_name} to port 8005 /instant-report."
)
doc.add_paragraph(
    "2. Alerts calls port 8000 for live BUY/SELL plus confidence."
)
doc.add_paragraph(
    "3. Alerts calls port 8002 for backtest confidence (optional)."
)
doc.add_paragraph("4. Builds HTML and sends via Brevo.")
doc.add_paragraph("5. The user receives the email.")
doc.add_paragraph(
    "The dashboard AI Signals page does the same ML path in-process (or via 8000) "
    "and draws charts."
)

doc.add_heading("10. What Is Not the Project", level=1)
doc.add_paragraph(
    "Do not study venv/Lib/site-packages/. That is pandas, Streamlit, yfinance, and other "
    "third-party libraries, not this application's source."
)

doc.add_heading("11. How to Start the Alerts Server", level=1)
doc.add_paragraph("From D:\\stock:")
doc.add_paragraph("python -m uvicorn alerts.main:app --host 0.0.0.0 --port 8005")
doc.add_paragraph("From D:\\stock\\alerts:")
doc.add_paragraph("python -m uvicorn main:app --host 0.0.0.0 --port 8005")
doc.add_paragraph("Then open http://localhost:8005/health. It should return active.")

out = Path(r"D:\stock\Code_Walkthrough.docx")
doc.save(str(out))
print(f"Wrote {out}")
